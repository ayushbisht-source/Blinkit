"""Geometry / text-fit QA for the deck.

LibreOffice cannot render in this sandbox, so overflow is estimated rather than seen.
Character-width factors are conservative (wider than reality) so the check errs toward
flagging, not toward passing something that would clip.
"""
import math
from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
MIN_FONT = 14

# average glyph width as a fraction of point size
WIDTH = {"Calibri": 0.47, "Cambria": 0.50, None: 0.50}

prs = Presentation("Blinkit-Category-Exploration.pptx")
issues = []
fonts_seen = set()

for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        try:
            x, y = Emu(sh.left).inches, Emu(sh.top).inches
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
        except TypeError:
            continue

        # out of bounds
        if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
            issues.append(f"S{idx}: shape off-canvas  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}")

        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text
        if not text.strip():
            continue

        # font size + face of the first run
        size = None
        face = None
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    size = r.font.size.pt
                if r.font.name:
                    face = r.font.name
                if size:
                    break
            if size:
                break
        if size is None:
            size = 18
        fonts_seen.add(face)

        if size < MIN_FONT:
            issues.append(f"S{idx}: font {size}pt below the {MIN_FONT}pt minimum — {text[:45]!r}")

        # margin check for text (shapes may bleed, text should not)
        if x < MARGIN - 0.01 or x + w > SLIDE_W - MARGIN + 0.01:
            issues.append(f"S{idx}: text box breaks the {MARGIN}\" margin — {text[:40]!r}")

        # estimated wrap
        cw = WIDTH.get(face, 0.50) * size / 72.0          # inches per char
        usable = w - 0.04                                  # margin:0 set on these boxes
        cpl = max(1, int(usable / cw))
        lines = 0
        for para in text.split("\n"):
            lines += max(1, math.ceil(len(para) / cpl))
        line_h = size * 1.22 / 72.0
        need = lines * line_h
        if need > h + 0.04:
            issues.append(
                f"S{idx}: est. overflow — needs {need:.2f}\" in {h:.2f}\" "
                f"({lines} lines @ {size}pt) — {text[:45]!r}"
            )

        boxes.append((x, y, w, h, text[:28]))

    # pairwise overlap of text boxes (shapes behind text are expected to overlap)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            ox = min(ax + aw, bx + bw) - max(ax, bx)
            oy = min(ay + ah, by + bh) - max(ay, by)
            if ox > 0.06 and oy > 0.06:
                issues.append(f"S{idx}: text overlap {at!r} x {bt!r} ({ox:.2f}\" x {oy:.2f}\")")

print(f"fonts used: {sorted(f for f in fonts_seen if f)}")
print(f"slides: {len(prs.slides)}")
if not issues:
    print("\nNo geometry or text-fit issues found.")
else:
    print(f"\n{len(issues)} issue(s):")
    for i in issues:
        print("  -", i)
